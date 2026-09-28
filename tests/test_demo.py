from __future__ import annotations

import json
import os
from pathlib import Path
import tomllib
import warnings

import ezdxf
from docx import Document
from fastapi.testclient import TestClient

from app.brain import BrainUnavailable, CodexBrainAdapter
from app.generators import generate_drawing, generate_presentation
from app.main import _disposable_model_path, _extract_brief_text, _read_site_boundary, _safe_readback, _sanitize_agent_reply, _validate_plan, create_app
from app.models import AgentSession, BuildPlan, DesignIR, EditPlan, ModelState, OutputManifest, Reference
from app.native_agent import AgentTurnResult, CodexAppServerRuntime
from app.references import FetchResult, ReferenceIngestor
from app.sketchup_mcp import SketchUpAdapter, _generated_script_dir
from app.store import ProjectStore
from tests.conftest import FakeBrain, FakeSketchUp, sample_context, sample_design, sample_plan


def test_schema_validation_and_plan_preflight():
    design = sample_design()
    plan = sample_plan()
    _validate_plan(design, plan)
    duplicate = design.model_dump(mode="json")
    duplicate["objects"].append(duplicate["objects"][1])
    try:
        DesignIR.model_validate(duplicate)
    except ValueError as error:
        assert "unique" in str(error)
    else:
        raise AssertionError("duplicate stable IDs should be rejected")


def test_codex_build_plan_is_compiled_from_design_ids():
    design = sample_design()
    plan = sample_plan().model_dump(mode="json")
    plan["operations"][0]["target_id"] = "HALLUCINATED_ID"
    plan["validation"]["required_ids"] = ["HALLUCINATED_ID"]
    plan["validation"]["expected_object_count_min"] = 99
    brain = CodexBrainAdapter(codex_executable="codex")
    brain._run = lambda *_args, **_kwargs: {  # type: ignore[method-assign]
        "design_ir": design.model_dump(mode="json"),
        "build_plan": plan,
        "decision_summary": "Three pavilions and a public passage.",
    }

    _, compiled, _ = brain.prepare(sample_context())

    geometry = [obj for obj in design.objects if obj.type in {"site_base", "building_mass", "circulation"}]
    geometry_operations = [op for op in compiled.operations if op.action in {"create_mass", "create_circulation"}]
    assert [op.target_id for op in geometry_operations] == [obj.id for obj in geometry]
    assert [op.action for op in geometry_operations] == [
        "create_circulation" if obj.type == "circulation" else "create_mass" for obj in geometry
    ]
    assert compiled.validation.required_ids == [obj.id for obj in geometry]
    assert compiled.validation.expected_object_count_min == len(geometry)
    assert [op.action for op in compiled.operations[-2:]] == ["capture_view", "save_model"]


def test_project_create_load_and_seed(tmp_path: Path):
    store = ProjectStore(tmp_path / "runtime")
    context = store.ensure_seed_project()
    assert context.project_id == "demo-cultural-center"
    loaded = store.load_project(context.project_id)
    assert loaded["context"].project_name == context.project_name
    assert loaded["model_state"].status == "ready"
    created = store.create_project(sample_context(""))
    assert created.project_id == "tidal-commons"
    assert store.load_project(created.project_id)["context"].project_id == created.project_id


def test_uploaded_brief_is_stored_in_project_context(tmp_path: Path):
    app = create_app(tmp_path / "runtime", brain=FakeBrain(), sketchup=FakeSketchUp())
    client = TestClient(app)
    client.get("/api/projects")
    upload = client.post(
        "/api/projects/demo-cultural-center/inputs/brief",
        params={"filename": "brief.txt"},
        content=b"Keep a clear route through the site.",
    )
    assert upload.status_code == 200
    path = upload.json()["path"]
    project = client.get("/api/projects/demo-cultural-center").json()
    assert path in project["context"]["brief"]["source_files"]
    assert (tmp_path / "runtime/projects/demo-cultural-center" / path).read_text(encoding="utf-8") == "Keep a clear route through the site."


def test_uploaded_reference_is_typed_and_serializes_without_warning(tmp_path: Path):
    app = create_app(tmp_path / "runtime", brain=FakeBrain(), sketchup=FakeSketchUp())
    client = TestClient(app)
    client.get("/api/projects")
    initial_reference_count = len(app.state.store.load_context("demo-cultural-center").references)

    with warnings.catch_warnings(record=True) as emitted:
        warnings.simplefilter("always")
        upload = client.post(
            "/api/projects/demo-cultural-center/inputs/reference",
            params={"filename": "jinshan-view.png"},
            content=b"reference-image-bytes",
        )

    assert upload.status_code == 200
    context = app.state.store.load_context("demo-cultural-center")
    assert len(context.references) == initial_reference_count + 1
    uploaded_reference = context.references[-1]
    assert isinstance(uploaded_reference, Reference)
    assert uploaded_reference.source == upload.json()["path"]
    assert not any("Pydantic serializer warnings" in str(item.message) for item in emitted)


def test_docx_brief_extraction_includes_table_content(tmp_path: Path):
    path = tmp_path / "table-taskbook.docx"
    document = Document()
    document.add_paragraph("Design taskbook")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Above-grade area"
    table.cell(0, 1).text = "20,900 m2"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Underground parking: 7,500 m2; do not omit from total area."
    document.save(path)

    extracted = _extract_brief_text(path)

    assert "Design taskbook" in extracted
    assert "Above-grade area" in extracted
    assert "20,900 m2" in extracted
    assert "Underground parking: 7,500 m2" in extracted
    assert extracted.count("Underground parking: 7,500 m2") == 1


def test_unknown_unit_site_dxf_prefers_redline_and_normalizes_survey_coordinates(tmp_path: Path):
    path = tmp_path / "site.dxf"
    document = ezdxf.new()
    document.units = 0
    modelspace = document.modelspace()
    modelspace.add_lwpolyline(
        [
            (462_552_000, 2_703_650_000),
            (462_652_000, 2_703_650_000),
            (462_652_000, 2_703_750_000),
            (462_552_000, 2_703_750_000),
            (462_552_000, 2_703_650_000),
        ],
        dxfattribs={"layer": "用地红线"},
    )
    modelspace.add_lwpolyline(
        [
            (462_400_000, 2_703_400_000),
            (462_760_000, 2_703_400_000),
            (462_760_000, 2_703_700_000),
            (462_400_000, 2_703_700_000),
        ],
        close=True,
        dxfattribs={"layer": "PUB_TITLE"},
    )
    document.saveas(path)

    result = _read_site_boundary(path)

    assert result is not None
    boundary, note = result
    assert boundary == [(0.0, 0.0), (100.0, 0.0), (100.0, 100.0), (0.0, 100.0)]
    assert "用地红线" in note
    assert "10000.000 m²" in note
    assert "millimeters inferred" in note


def test_dxf_and_presentation_generation(tmp_path: Path):
    context = sample_context()
    design = sample_design()
    project_dir = tmp_path / "project"
    project_dir.mkdir()
    dxf_path, svg_path = generate_drawing(project_dir, design)
    document = ezdxf.readfile(dxf_path)
    modelspace = document.modelspace()
    assert len(modelspace.query("LWPOLYLINE")) >= 5
    assert {layer.dxf.name for layer in document.layers} >= {"SITE", "MASS", "CIRCULATION", "LABELS"}
    assert "PUBLIC PASSAGE" in svg_path.read_text(encoding="utf-8").upper()
    board = generate_presentation(project_dir, context, design, svg_path, None)
    text = board.read_text(encoding="utf-8")
    assert "A3 landscape" in text
    assert "Reading Pavilion" in text
    assert "data:image/png;base64" not in text


def test_sketchup_adapter_wraps_existing_tools():
    class RecordingClient:
        def __init__(self):
            self.calls = []

        def call(self, name, args=None):
            self.calls.append((name, args or {}))
            return {"entity_id": 712}

    client = RecordingClient()
    adapter = SketchUpAdapter(client)
    result = adapter.create_mass(stable_id="MASS_01", name="Reading Pavilion", origin_m=[1, 2, 0], width_m=12, depth_m=8, height_m=7.2)
    assert result["entity_id"] == 712
    assert client.calls[0][0] == "sketchup_create_mass"
    assert client.calls[0][1]["name"] == "[MASS_01] Reading Pavilion"
    adapter.modify_object(connector_ref="712", translate_m=[3, 0, 0])
    assert client.calls[1][0] == "sketchup_transform_group"
    assert client.calls[1][1]["entity_id"] == 712


def test_connector_readback_preserves_geometry_bounds():
    result = _safe_readback({"entity_id": 712, "bounds_m": {"width": 15.0, "depth": 8.0}, "model_bounds_m": {"width": 60.0}})
    assert result == {"bounds_m": {"width": 15.0, "depth": 8.0}, "model_bounds_m": {"width": 60.0}}


def test_sketchup_adapter_saves_a_copy_without_changing_active_model_path(tmp_path: Path, monkeypatch):
    generated_scripts = tmp_path / "generated_scripts"
    monkeypatch.setenv("ARCHFLOW_GENERATED_SCRIPT_DIR", str(generated_scripts))

    class SaveRecordingClient:
        script = ""

        def call(self, name, args=None):
            assert name == "sketchup_eval_project_file"
            self.script = Path(args["script_path"]).read_text(encoding="utf-8")
            return {"saved": True}

    client = SaveRecordingClient()
    target = tmp_path / "outputs" / "demo.skp"
    SketchUpAdapter(client).save_model(target)
    assert "model.save_copy(" in client.script
    assert "model.save_as(" not in client.script


def test_brain_job_lifecycle_and_project_prepare(tmp_path: Path):
    class OfflineBrain(FakeBrain):
        def prepare(self, context, images=None, previous_design=None):
            raise BrainUnavailable("Codex CLI is not available.")

    app = create_app(tmp_path / "runtime", brain=OfflineBrain(), sketchup=FakeSketchUp())
    client = TestClient(app)
    client.get("/api/projects")
    response = client.post("/api/projects/demo-cultural-center/prepare", json={"user_intent": "Keep the route open."})
    assert response.status_code == 200
    job_id = response.json()["job"]["job_id"]
    assert response.json()["status"] == "awaiting_codex"
    completion = client.post(f"/api/jobs/{job_id}/complete", json={
        "design_ir": sample_design().model_dump(mode="json"),
        "build_plan": sample_plan().model_dump(mode="json"),
        "decision_summary": "Three pavilions with a public route.",
    })
    assert completion.status_code == 200
    saved = client.get("/api/projects/demo-cultural-center")
    assert saved.json()["design_ir"]["objects"][0]["id"] == "SITE_BASE_01"
    assert Path(tmp_path / "runtime/jobs" / job_id / "job.json").exists()


def test_build_and_two_edits_use_the_same_model(tmp_path: Path):
    brain, sketchup = FakeBrain(), FakeSketchUp()
    app = create_app(tmp_path / "runtime", brain=brain, sketchup=sketchup)
    client = TestClient(app)
    client.get("/api/projects")
    prepared = client.post("/api/projects/demo-cultural-center/prepare", json={})
    assert prepared.status_code == 200
    built = client.post("/api/projects/demo-cultural-center/build", json={"confirm_disposable_model": True})
    assert built.status_code == 200
    assert len(built.json()["model_state"]["objects"]) == 5
    first = client.post("/api/projects/demo-cultural-center/edit", json={"instruction": "Raise the reading pavilion by one floor."})
    second = client.post("/api/projects/demo-cultural-center/edit", json={"instruction": "Move the learning house three meters east."})
    assert first.status_code == second.status_code == 200
    calls = [call for call in sketchup.calls if call[0] == "modify_object"]
    assert len(calls) == 2
    assert calls[0][1]["connector_ref"] != calls[1][1]["connector_ref"]
    final_project = client.get("/api/projects/demo-cultural-center").json()
    model = final_project["model_state"]
    assert model["status"] == "edited"
    assert len([item for item in model["objects"] if item.get("last_change")]) == 2
    route = next(item for item in model["objects"] if item["stable_id"] == "PUBLIC_STREET_01")
    assert route["origin"][2] == 0.2
    road_call = next(call for call in sketchup.calls if call[0] == "create_circulation")
    assert road_call[1]["start_m"][2] == road_call[1]["end_m"][2] == 0.2
    shifted_workshop = next(item for item in final_project["design_ir"]["objects"] if item["id"] == "MASS_02")
    assert shifted_workshop["footprint"] == [[27, 8], [39, 8], [39, 18], [27, 18]]
    assert len(final_project["output_manifest"]["render"]) == 3
    assert final_project["output_manifest"]["drawing"]
    assert final_project["output_manifest"]["presentation"]
    assert any(call[0] == "get_model_info" for call in sketchup.calls)


def test_partial_build_can_resume_without_duplicate_geometry(tmp_path: Path):
    class SaveFailsOnceSketchUp(FakeSketchUp):
        save_attempts = 0

        def save_model(self, target_path: Path, operation_name: str = "checkpoint"):
            self.save_attempts += 1
            if self.save_attempts == 1:
                raise ValueError("Simulated first-save failure.")
            return super().save_model(target_path, operation_name)

    sketchup = SaveFailsOnceSketchUp()
    app = create_app(tmp_path / "runtime", brain=FakeBrain(), sketchup=sketchup)
    client = TestClient(app)
    client.get("/api/projects")
    assert client.post("/api/projects/demo-cultural-center/prepare", json={}).status_code == 200
    first = client.post("/api/projects/demo-cultural-center/build", json={"confirm_disposable_model": True})
    assert first.status_code == 502
    partial = client.get("/api/projects/demo-cultural-center").json()["model_state"]
    assert partial["status"] == "partial"
    assert len(partial["objects"]) == 5

    resumed = client.post("/api/projects/demo-cultural-center/build", json={"confirm_disposable_model": True})
    assert resumed.status_code == 200
    assert resumed.json()["status"] == "built"
    create_calls = [call for call in sketchup.calls if call[0] in {"create_mass", "create_circulation"}]
    assert len(create_calls) == 5
    assert sketchup.save_attempts == 2


def test_reference_ingestor_blocks_private_network_and_extracts_visible_text(monkeypatch):
    blocked = ReferenceIngestor().ingest("http://127.0.0.1/private")
    assert blocked["status"] == "unreadable"
    assert "不允许读取本机" in blocked["error"]

    def private_dns(_host, _port, type=None):
        return [(2, 1, 6, "", ("10.0.0.7", 80))]

    monkeypatch.setattr("app.references.socket.getaddrinfo", private_dns)
    private = ReferenceIngestor().ingest("https://studio.example/design")
    assert private["status"] == "unreadable"
    assert "不允许读取本机" in private["error"]

    class ReadablePage(ReferenceIngestor):
        def _validate_and_resolve(self, raw_url):
            from urllib.parse import urlsplit
            return urlsplit(raw_url), ["93.184.216.34"], 443

        def _request_once(self, parts, port, address, path):
            body = b"<html><head><title>Waterfront Library</title><script>ignore me</script></head><body><nav>menu text</nav><main><h1>Public route and shaded courtyard</h1><p>Readable architectural precedent with an open public passage.</p></main></body></html>"
            return FetchResult(200, "text/html; charset=utf-8", body)

    readable = ReadablePage().ingest("https://studio.example/design")
    assert readable["status"] == "readable"
    assert readable["title"] == "Waterfront Library"
    assert "Public route" in readable["excerpt"]
    assert "ignore me" not in readable["excerpt"]
    assert "menu text" not in readable["excerpt"]

    class RedirectToLocal(ReferenceIngestor):
        requests = 0

        def _validate_and_resolve(self, raw_url):
            from urllib.parse import urlsplit
            if self.requests:
                return super()._validate_and_resolve(raw_url)
            return urlsplit(raw_url), ["93.184.216.34"], 443

        def _request_once(self, parts, port, address, path):
            self.requests += 1
            return FetchResult(302, "text/html", b"", "http://127.0.0.1/private")

    redirected = RedirectToLocal().ingest("https://studio.example/design")
    assert redirected["status"] == "unreadable"
    assert "不允许读取本机" in redirected["error"]


def test_reference_fetch_enforces_response_size_cap(monkeypatch):
    class FakeResponse:
        status = 200
        def getheader(self, key, default=None):
            return {"Content-Length": "11", "Content-Type": "text/plain"}.get(key, default)
    class FakeConnection:
        sock = None
        def __init__(self, *args):
            pass
        def request(self, *args, **kwargs):
            pass
        def getresponse(self):
            return FakeResponse()
        def close(self):
            pass

    monkeypatch.setattr("app.references._PinnedHTTPConnection", FakeConnection)
    from urllib.parse import urlsplit
    try:
        ReferenceIngestor(max_bytes=10)._request_once(urlsplit("http://example.com/"), 80, "93.184.216.34", "/")
    except ValueError as error:
        assert "1 MB" in str(error)
    else:
        raise AssertionError("oversized pages should be rejected")


def test_prepare_persists_reference_ingestion_and_screenshot_fallback(tmp_path: Path):
    app = create_app(tmp_path / "runtime", brain=FakeBrain(), sketchup=FakeSketchUp())
    client = TestClient(app)
    client.get("/api/projects")
    response = client.post("/api/projects/demo-cultural-center/prepare", json={"reference_url": "http://localhost:45678/case"})
    assert response.status_code == 200
    assert response.json()["reference_warnings"]
    project = client.get("/api/projects/demo-cultural-center").json()
    reference = next(item for item in project["context"]["references"] if item["type"] == "url")
    assert reference["status"] == "unreadable"
    assert "上传网页截图或参考图片" in reference["error"]


def test_chinese_conversation_refines_design_before_build(tmp_path: Path):
    brain = FakeBrain()
    native = FakeNativeAgent()
    app = create_app(tmp_path / "runtime", brain=brain, sketchup=FakeSketchUp(), native_agent=native)
    client = TestClient(app)
    client.get("/api/projects")
    first = client.post("/api/projects/demo-cultural-center/conversation", json={
        "message": "公共街道再宽一点，两个主要体块之间更开放。",
        "project_name": "潮汐共享 · 滨水文化之家",
        "brief": "为滨水社区设计公共文化空间。",
        "site_note": "60 m × 48 m 合成场地。",
        "user_intent": "保持低矮、开放。",
    })
    assert first.status_code == 200
    assert first.json()["project"]["design_ir"] is None
    assert native.calls[-1]["mcp_enabled"] is False

    second = client.post("/api/projects/demo-cultural-center/conversation", json={"message": "公共街道再加宽一些。"})
    assert second.status_code == 200
    assert native.calls[-1]["mcp_enabled"] is False
    assert native.calls[-1]["thread_id"] == "thr-fast-assembly"
    assert not brain.prepare_calls
    messages = second.json()["project"]["context"]["conversation"]
    assert [item["role"] for item in messages] == ["user", "assistant", "user", "assistant"]


def test_chinese_conversation_edits_same_built_model_and_width_depth_routes(tmp_path: Path):
    brain, sketchup = FakeBrain(), FakeSketchUp()
    app = create_app(tmp_path / "runtime", brain=brain, sketchup=sketchup)
    client = TestClient(app)
    client.get("/api/projects")
    assert client.post("/api/projects/demo-cultural-center/prepare", json={}).status_code == 200
    assert client.post("/api/projects/demo-cultural-center/build", json={"confirm_disposable_model": True}).status_code == 200

    width = client.post("/api/projects/demo-cultural-center/edit", json={"instruction": "把阅览体块宽度扩大到 15 米。"})
    depth = client.post("/api/projects/demo-cultural-center/edit", json={"instruction": "把阅览体块进深改为 14 米。"})
    assert width.status_code == depth.status_code == 200
    mass = next(item for item in depth.json()["project"]["model_state"]["objects"] if item["stable_id"] == "MASS_01")
    assert mass["width"] == 15.0
    assert mass["depth"] == 14.0
    assert len([call for call in sketchup.calls if call[0] == "modify_object"]) == 2


def test_invalid_dimension_edit_is_rejected_before_connector_call(tmp_path: Path):
    class InvalidWidthBrain(FakeBrain):
        def plan_edit(self, context, design, model_state, instruction):
            return EditPlan(target_id="MASS_01", patch={"width": 0.2}, rationale="invalid test")

    sketchup = FakeSketchUp()
    app = create_app(tmp_path / "runtime", brain=InvalidWidthBrain(), sketchup=sketchup)
    client = TestClient(app)
    client.get("/api/projects")
    assert client.post("/api/projects/demo-cultural-center/prepare", json={}).status_code == 200
    assert client.post("/api/projects/demo-cultural-center/build", json={"confirm_disposable_model": True}).status_code == 200
    response = client.post("/api/projects/demo-cultural-center/edit", json={"instruction": "扩大体块宽度。"})
    assert response.status_code == 502
    assert not [call for call in sketchup.calls if call[0] == "modify_object"]


def test_workspace_copy_is_simplified_chinese_and_has_one_conversation_box(tmp_path: Path):
    app = create_app(tmp_path / "runtime", brain=FakeBrain(), sketchup=FakeSketchUp())
    page = TestClient(app).get("/")
    assert page.status_code == 200
    assert 'lang="zh-CN"' in page.text
    assert "讨论与修改" in page.text
    assert "发送消息" in page.text
    assert "旧版规则化建模流程" in page.text
    assert "自由建模会话" in page.text
    assert "MODEL ROUTER V1" in page.text


class FakeNativeAgent:
    available = True
    model = "gpt-6-sol"
    reasoning_effort = "medium"
    failed_tool_calls = 0

    def __init__(self):
        self.calls = []

    def respond(self, **kwargs):
        self.calls.append(kwargs)
        return AgentTurnResult(
            thread_id=kwargs.get("thread_id") or "thr-fast-assembly",
            reply="已根据要求完成建筑设计推演。" if kwargs["mcp_enabled"] else "已记录设计方向。",
            tool_calls=([{"server": "kongxing_sketchup", "tool": "sketchup_create_gable_roof"},
                         {"server": "kongxing_sketchup", "tool": "sketchup_export_view_image"}]
                        if kwargs["mcp_enabled"] else []),
            model_name=kwargs.get("model", self.model),
            failed_tool_calls=self.failed_tool_calls,
        )


def test_local_status_exposes_native_model_and_reasoning_effort(tmp_path: Path):
    native = FakeNativeAgent()
    app = create_app(tmp_path / "runtime", brain=FakeBrain(), sketchup=FakeSketchUp(), native_agent=native)
    response = TestClient(app).get("/api/status")
    assert response.status_code == 200
    assert response.json()["native_agent_model"] == "gpt-6-sol"
    assert response.json()["native_agent_reasoning_effort"] == "medium"


class DisconnectedFakeSketchUp(FakeSketchUp):
    def ping(self):
        if not self.active_model_path:
            from app.sketchup_mcp import ConnectorUnavailable
            raise ConnectorUnavailable("SketchUp bridge is not running.")
        return super().ping()


def test_native_conversation_gates_sketchup_tools_until_same_disposable_model(tmp_path: Path):
    runtime = tmp_path / "runtime"
    sketchup = DisconnectedFakeSketchUp()
    native = FakeNativeAgent()
    blank_path = runtime / "projects" / "demo-cultural-center" / "outputs" / "model" / "blank-disposable-test.skp"

    def launch(project_id, runtime_root, existing_model_path=None):
        assert project_id == "demo-cultural-center"
        assert existing_model_path is None
        blank_path.parent.mkdir(parents=True, exist_ok=True)
        blank_path.write_bytes(b"disposable copy")
        sketchup.active_model_path = str(blank_path)
        return blank_path

    app = create_app(runtime, brain=FakeBrain(), sketchup=sketchup, native_agent=native,
                     disposable_model_launcher=launch)
    client = TestClient(app)
    client.get("/api/projects")

    discussion = client.post("/api/projects/demo-cultural-center/conversation", json={"message": "入口希望面向东侧。"})
    assert discussion.status_code == 200
    assert native.calls[-1]["mcp_enabled"] is False
    assert not [call for call in sketchup.calls if call[0] in {"create_mass", "create_circulation", "modify_object"}]

    started = client.post("/api/projects/demo-cultural-center/agent/session", json={"confirm_disposable_model": True})
    assert started.status_code == 200
    assert started.json()["session"]["status"] == "ready"
    assert started.json()["session"]["model_path"].endswith("blank-disposable-test.skp")

    modeled = client.post("/api/projects/demo-cultural-center/conversation", json={"message": "增加一个坡顶，并把公共连廊接到场地入口。"})
    assert modeled.status_code == 200
    assert native.calls[-1]["mcp_enabled"] is True
    assert native.calls[-1]["thread_id"] == "thr-fast-assembly"
    assert len(modeled.json()["agent"]["tool_calls"]) == 2
    project = modeled.json()["project"]
    assert project["agent_session"]["status"] == "ready"
    assert project["agent_session"]["thread_id"] == "thr-fast-assembly"
    assert project["model_state"]["model_path"].endswith("fast-assembly-agent.skp")
    assert any(item["type"] == "viewport" for item in project["output_manifest"]["render"])
    assert [item["phase"] for item in project["context"]["conversation"]] == ["agent", "agent", "agent", "agent"]
    premium = client.post(
        "/api/projects/demo-cultural-center/conversation",
        json={"message": "请精修一次入口和公共空间的模型表现。", "tier": "premium"},
    )
    assert premium.status_code == 200
    assert premium.json()["agent"]["model"] == "gpt-6-astra"
    assert premium.json()["agent"]["tier"] == "premium"
    assert premium.json()["agent"]["reasoning_effort"] == "low"
    routine = client.post(
        "/api/projects/demo-cultural-center/conversation",
        json={"message": "读取模型状态并简要报告。"},
    )
    assert routine.status_code == 200
    assert routine.json()["agent"]["model"] == "gpt-6-sol"
    assert routine.json()["agent"]["tier"] == "economy"


def test_economy_failures_suggest_premium_but_never_call_astra_without_selection(tmp_path: Path):
    class FailingToolNative(FakeNativeAgent):
        failed_tool_calls = 1

    app = create_app(tmp_path / "runtime", brain=FakeBrain(), sketchup=FakeSketchUp(), native_agent=FailingToolNative())
    client = TestClient(app)
    client.get("/api/projects")
    project_path = "/api/projects/demo-cultural-center/conversation"

    first = client.post(project_path, json={"message": "先讨论入口。", "tier": "economy"})
    second = client.post(project_path, json={"message": "再讨论路径。", "tier": "economy"})
    assert first.json()["agent"]["model"] == second.json()["agent"]["model"] == "gpt-6-sol"
    assert second.json()["agent"]["premium_rescue_pending"] is True

    continued = client.post(project_path, json={"message": "继续检查失败的工具调用。", "tier": "economy"})
    assert continued.status_code == 200
    assert continued.json()["agent"]["model"] == "gpt-6-sol"
    assert continued.json()["agent"]["tier"] == "economy"
    assert continued.json()["agent"]["premium_rescue_pending"] is True

    premium = client.post(project_path, json={"message": "我选择精修。", "tier": "premium"})
    assert premium.status_code == 200
    assert premium.json()["agent"]["model"] == "gpt-6-astra"
    assert premium.json()["agent"]["premium_rescue_pending"] is False

    followup = client.post(project_path, json={"message": "恢复普通迭代。", "tier": "economy"})
    assert followup.status_code == 200
    assert followup.json()["agent"]["model"] == "gpt-6-sol"


def test_native_agent_refuses_original_or_switched_sketchup_model(tmp_path: Path):
    runtime = tmp_path / "runtime"
    sketchup = FakeSketchUp()
    native = FakeNativeAgent()
    app = create_app(runtime, brain=FakeBrain(), sketchup=sketchup, native_agent=native)
    client = TestClient(app)
    client.get("/api/projects")
    project_dir = runtime / "projects" / "demo-cultural-center"
    blank_path = project_dir / "outputs" / "model" / "blank-disposable-test.skp"
    blank_path.parent.mkdir(parents=True, exist_ok=True)
    blank_path.write_bytes(b"disposable copy")
    store = ProjectStore(runtime)
    session = AgentSession(project_id="demo-cultural-center", status="ready", thread_id="thr-existing", model_path="outputs/model/blank-disposable-test.skp")
    store.save_state("demo-cultural-center", session, "agent_session.json")
    sketchup.active_model_path = str(tmp_path / "private-original.skp")

    response = client.post("/api/projects/demo-cultural-center/conversation", json={"message": "移动二层体块。"})
    assert response.status_code == 409
    assert not native.calls
    assert not [call for call in sketchup.calls if call[0] == "modify_object"]


def test_native_runtime_isolates_mcp_allowlist_and_uses_codex_login_cache(tmp_path: Path, monkeypatch):
    source_home = tmp_path / "user-codex"
    source_home.mkdir()
    (source_home / "config.toml").write_text(
        'model = "gpt-6-luna"\n'
        '[mcp_servers.kongxing_sketchup]\ncommand = "node"\nargs = ["server.mjs"]\n'
        '[mcp_servers.honglu-autocad]\nurl = "http://127.0.0.1:9000/mcp"\n',
        encoding="utf-8",
    )
    (source_home / "auth.json").write_text("opaque test login cache", encoding="utf-8")
    monkeypatch.setenv("CODEX_HOME", str(source_home))
    class ToolClient:
        def list_tools(self):
            return [
                {"name": "sketchup_health", "description": "Read bridge health.", "inputSchema": {"type": "object", "properties": {}}},
                {"name": "sketchup_create_gable_roof", "description": "Create a gable roof.", "inputSchema": {"type": "object", "properties": {"rise_m": {"type": "number"}}}},
            ]

    runtime = CodexAppServerRuntime(tmp_path / "runtime", codex_executable="codex-test",
                                   sketchup_mcp=ToolClient(), home_root=tmp_path / "codex-home")

    runtime._prepare_home(mcp_enabled=True)
    isolated = tomllib.loads((runtime.home / "config.toml").read_text(encoding="utf-8"))
    assert isolated["model"] == "gpt-6-sol"
    assert "mcp_servers" not in isolated
    assert os.path.samefile(source_home / "auth.json", runtime.home / "auth.json")
    dynamic_tools = runtime._dynamic_tools()
    assert [tool["name"] for tool in dynamic_tools] == ["sketchup_health", "sketchup_create_gable_roof"]
    assert dynamic_tools[1]["inputSchema"]["properties"]["rise_m"]["type"] == "number"

    runtime._prepare_home(mcp_enabled=False)
    no_tools = tomllib.loads((runtime.home / "config.toml").read_text(encoding="utf-8"))
    assert "mcp_servers" not in no_tools


def test_disposable_model_identity_check_allows_only_project_runtime_copy(tmp_path: Path):
    project_dir = tmp_path / "runtime" / "projects" / "demo-cultural-center"
    allowed = project_dir / "outputs" / "model" / "blank-disposable-demo.skp"
    allowed.parent.mkdir(parents=True)
    allowed.write_bytes(b"blank")
    assert _disposable_model_path(str(allowed), project_dir) == allowed.resolve()
    assert _disposable_model_path(str(tmp_path / "private-thesis.skp"), project_dir) is None
    assert _disposable_model_path(str(project_dir / "outputs/model/private.skp"), project_dir) is None


def test_agent_reply_hides_host_local_paths():
    reply = "查看截图：[完成截图](R:/connector-runtime/screenshots/view.png)，记录在 D:\\demo\\model.skp"
    sanitized = _sanitize_agent_reply(reply)
    assert "R:/connector-runtime" not in sanitized
    assert "D:\\demo" not in sanitized
    assert "完成截图（本机路径已隐藏）" in sanitized
    assert "[本机路径已隐藏]" in sanitized


def test_generated_sketchup_script_dir_uses_local_environment(tmp_path: Path, monkeypatch):
    monkeypatch.delenv("ARCHFLOW_GENERATED_SCRIPT_DIR", raising=False)
    monkeypatch.setenv("PROGRAMDATA", str(tmp_path / "program-data"))
    assert _generated_script_dir() == tmp_path / "program-data" / "archflow-mcp" / "generated_scripts"

    override = tmp_path / "connector-scripts"
    monkeypatch.setenv("ARCHFLOW_GENERATED_SCRIPT_DIR", str(override))
    assert _generated_script_dir() == override
